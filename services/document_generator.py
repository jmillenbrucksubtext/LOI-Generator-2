"""
Document generator that opens a Word template, handles conditional scenarios,
replaces placeholders with tracked changes, and inserts property photos.

Uses python-docx for document structure and lxml for tracked-change XML
(python-docx has no native tracked-changes API).
"""

import copy
import io
import os
import struct
from datetime import datetime
from typing import Optional

from docx import Document
from docx.shared import Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

from services.loi_form_data import (
    LoiFormData,
    DepositStructure,
    DueDiligenceType,
    CommissionType,
    SignatureBlockType,
)
from services.number_to_words import to_legal_dollar_string, convert_to_words

# OpenXML namespaces
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
PIC_NS = "http://schemas.openxmlformats.org/drawingml/2006/picture"

NSMAP = {
    "w": W_NS,
    "r": R_NS,
    "a": A_NS,
    "wp": WP_NS,
    "pic": PIC_NS,
}


def _qn(tag: str) -> str:
    """Convert a namespace-prefixed tag like 'w:r' to Clark notation."""
    prefix, local = tag.split(":")
    return f"{{{NSMAP[prefix]}}}{local}"


class DocumentGenerator:
    DEFAULT_AUTHOR = "Subtext LOI Generator"

    def __init__(self):
        self._author = self.DEFAULT_AUTHOR
        self._revision_id = 1

    def _next_rev_id(self) -> str:
        rid = str(self._revision_id)
        self._revision_id += 1
        return rid

    # ------------------------------------------------------------------
    # Public entry point
    # ------------------------------------------------------------------
    def generate(self, template_path: str, form: LoiFormData) -> io.BytesIO:
        full_name = f"{form.prepared_by_first_name} {form.prepared_by_last_name}".strip()
        self._author = full_name if full_name else self.DEFAULT_AUTHOR

        doc = Document(template_path)
        body = doc.element.body
        now = datetime.now().strftime("%Y-%m-%dT%H:%M:%SZ")

        # Step 1: Handle scenario sections
        self._handle_scenarios(body, form, now)

        # Step 2: Replace bracketed placeholders (body)
        self._replace_placeholders(body, form, now)

        # Step 3: Replace placeholders in headers
        self._replace_header_placeholders(doc, form, now)

        # Step 4: Insert property photo
        if form.property_photo_bytes and len(form.property_photo_bytes) > 0:
            self._insert_property_photo(doc, body, form, now)

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf

    # ------------------------------------------------------------------
    # Scenario Handling
    # ------------------------------------------------------------------
    def _handle_scenarios(self, body, form: LoiFormData, now: str):
        self._handle_deposit_scenario(body, form, now)
        self._handle_legal_reimbursement(body, form, now)
        self._handle_due_diligence_scenario(body, form, now)
        self._handle_closing_extension(body, form, now)
        self._handle_commission_scenario(body, form, now)
        self._handle_option_to_extend(body, form, now)
        self._handle_lease_scenario(body, form, now)
        self._handle_seller_rollover(body, form, now)
        self._handle_signature_block(body, form, now)
        self._handle_multiple_entities(body, form, now)

    def _handle_deposit_scenario(self, body, form: LoiFormData, now: str):
        marker1 = "REMOVE PARAGRAPH ABOVE AND USE THE FOLLOWING PARAGRPAPH IF THE INITIAL DEPOSIT IS GOING HARD AT THE END OF DUE DILIGENCE."
        marker2 = "USE THE FOLLOWING PARAGRAPH IF MONEY IS GOING HARD MONTHLY AFTER DUE DILIGENCE."

        for para in list(body.iterchildren(_qn("w:p"))):
            text = _get_paragraph_text(para)
            if "Initial Deposit" not in text or marker1 not in text:
                continue

            part1_end = text.index(marker1)
            part2_start = part1_end + len(marker1)
            part2_end = text.index(marker2)
            part3_start = part2_end + len(marker2)

            if form.deposit_structure == DepositStructure.DUE_DILIGENCE_GOING_HARD:
                segments = [
                    (0, part2_start, True),
                    (part2_start, part2_end - part2_start, False),
                    (part2_end, len(text) - part2_end, True),
                ]
            elif form.deposit_structure == DepositStructure.MONTHLY_GOING_HARD:
                segments = [
                    (0, part3_start, True),
                    (part3_start, len(text) - part3_start, False),
                ]
            else:  # GovernmentalApprovalsGoingHard
                segments = [
                    (0, part1_end, False),
                    (part1_end, len(text) - part1_end, True),
                ]

            self._rebuild_paragraph_with_scenario(para, text, segments, now)
            return

    def _handle_legal_reimbursement(self, body, form: LoiFormData, now: str):
        marker = "INSERT THIS PARAGRAPH IF WE ARE PAYING A LEGAL REIMBURSEMENT FEE AT PSA EXECUTION"
        self._handle_optional_paragraph(body, marker, form.include_legal_reimbursement, now)

    def _handle_due_diligence_scenario(self, body, form: LoiFormData, now: str):
        marker = "REPLACE THIS ENTIRE PARAGRAPH WITH THE FOLLOWING WHEN REQUIRING AN ASSEMBLAGE PERIOD"
        for para in list(body.iterchildren(_qn("w:p"))):
            text = _get_paragraph_text(para)
            if marker not in text:
                continue

            marker_idx = text.index(marker)
            if form.due_diligence_type == DueDiligenceType.STANDARD:
                segments = [
                    (0, marker_idx, False),
                    (marker_idx, len(text) - marker_idx, True),
                ]
            else:
                after_marker = marker_idx + len(marker)
                segments = [
                    (0, after_marker, True),
                    (after_marker, len(text) - after_marker, False),
                ]
            self._rebuild_paragraph_with_scenario(para, text, segments, now)
            return

    def _handle_closing_extension(self, body, form: LoiFormData, now: str):
        marker = "INSERT THE FOLLOWING LANGUAGE IF WE ARE INCLUDING A CLOSING EXTENSION"
        for para in list(body.iterchildren(_qn("w:p"))):
            text = _get_paragraph_text(para)
            if marker not in text:
                continue

            marker_idx = text.index(marker)
            if form.include_closing_extension:
                segments = [
                    (0, marker_idx, False),
                    (marker_idx, len(marker), True),
                    (marker_idx + len(marker), len(text) - marker_idx - len(marker), False),
                ]
            else:
                segments = [
                    (0, marker_idx, False),
                    (marker_idx, len(text) - marker_idx, True),
                ]
            self._rebuild_paragraph_with_scenario(para, text, segments, now)
            return

    def _handle_commission_scenario(self, body, form: LoiFormData, now: str):
        marker_seller = "IF SELLER IS PAYING COMMISSION TO LISTING AGENT"
        marker_subtext = "IF SUBTEXT IS PAYING THE COMMISSION"
        marker_no_brokers = "IF NO BROKERS ARE INVOLVED IN THIS TRANSACTION"

        for para in list(body.iterchildren(_qn("w:p"))):
            text = _get_paragraph_text(para)
            if marker_seller not in text or "Commissions." not in text:
                continue

            idx1 = text.index(marker_seller)
            after1 = idx1 + len(marker_seller)
            idx2 = text.index(marker_subtext)
            after2 = idx2 + len(marker_subtext)
            idx3 = text.index(marker_no_brokers)
            after3 = idx3 + len(marker_no_brokers)

            prefix_end = text.index("Commissions.") + len("Commissions.")
            segments = [(0, prefix_end, False)]

            if form.commission_type == CommissionType.SELLER_PAYS_LISTING_AGENT:
                segments.append((prefix_end, idx1 - prefix_end, True))
                segments.append((idx1, len(marker_seller), True))
                segments.append((after1, idx2 - after1, False))
                segments.append((idx2, len(text) - idx2, True))
            elif form.commission_type == CommissionType.SUBTEXT_PAYS:
                segments.append((prefix_end, idx2 - prefix_end, True))
                segments.append((idx2, len(marker_subtext), True))
                segments.append((after2, idx3 - after2, False))
                segments.append((idx3, len(text) - idx3, True))
            else:  # NoBrokers
                segments.append((prefix_end, idx3 - prefix_end, True))
                segments.append((idx3, len(marker_no_brokers), True))
                segments.append((after3, len(text) - after3, False))

            self._rebuild_paragraph_with_scenario(para, text, segments, now)
            return

    def _handle_option_to_extend(self, body, form: LoiFormData, now: str):
        if not form.include_option_to_extend:
            for para in list(body.iterchildren(_qn("w:p"))):
                text = _get_paragraph_text(para)
                if "Option to Extend." in text and "Extension Notice" in text:
                    self._delete_entire_paragraph(para, now)
                    return
            return
        # Replace number of extensions and days per extension within the paragraph
        for para in list(body.iterchildren(_qn("w:p"))):
            text = _get_paragraph_text(para)
            if "Option to Extend." not in text or "Extension Notice" not in text:
                continue
            if form.num_extension_options is not None and form.num_extension_options != 2:
                self._replace_text_in_paragraph(
                    para, "two (2)", _format_period(form.num_extension_options), now)
            if form.extension_option_days is not None and form.extension_option_days != 60:
                self._replace_text_in_paragraph(
                    para, "sixty (60)", _format_period(form.extension_option_days), now)
            return

    def _handle_lease_scenario(self, body, form: LoiFormData, now: str):
        marker_vacant = "REMOVE THE PRIOR SENTENCE AND USE THE FOLLOWING IF BEING DELIVERD VACANT"
        marker_termination = "REMOVE THE PRIOR SENTENCE AND ADD THE FOLLOWING IF LEASES WILL EXIST AND LANDLORD IS NOT TERMINATING THEM"
        marker_negotiate = "USE THE FOLLOWING IF WE NEED THE RIGHT TO NEGOTIATE WITH EXISTING TENANTS PRIOR TO CLOSING."

        for para in list(body.iterchildren(_qn("w:p"))):
            text = _get_paragraph_text(para)
            if "Leases." not in text or marker_vacant not in text:
                continue

            idx_vacant = text.index(marker_vacant)
            idx_termination = text.index(marker_termination)
            idx_negotiate = text.index(marker_negotiate)

            lease_start = text.index("Leases.")
            prefix_end = lease_start + len("Leases.")

            vacant_end = idx_vacant + len(marker_vacant)
            termination_end = idx_termination + len(marker_termination)
            negotiate_end = idx_negotiate + len(marker_negotiate)

            segments = []
            segments.append((0, prefix_end, False))
            segments.append((prefix_end, idx_vacant - prefix_end, not form.include_existing_leases))
            segments.append((idx_vacant, len(marker_vacant), True))
            segments.append((vacant_end, idx_termination - vacant_end, not form.include_delivered_vacant))
            segments.append((idx_termination, len(marker_termination), True))
            segments.append((termination_end, idx_negotiate - termination_end, not form.include_lease_termination))

            if form.include_right_to_negotiate_with_tenants:
                segments.append((idx_negotiate, len(marker_negotiate), True))
                segments.append((negotiate_end, len(text) - negotiate_end, False))
            else:
                segments.append((idx_negotiate, len(text) - idx_negotiate, True))

            self._rebuild_paragraph_with_scenario(para, text, segments, now)
            return

    def _handle_seller_rollover(self, body, form: LoiFormData, now: str):
        marker = "INSERT THIS PARAGRAGH IF WE ARE ALLOWING THE SELLER TO CONTRIBUTE LAND AND RECEIVE LP LEVEL RETURNS"
        # Template has marker in its own paragraph, content in the next paragraph.
        # Both have numId=0 (numbering disabled).
        paragraphs = list(body.iterchildren(_qn("w:p")))
        for i, para in enumerate(paragraphs):
            text = _get_paragraph_text(para)
            if marker not in text:
                continue
            # Always delete the marker paragraph
            self._delete_entire_paragraph(para, now)
            if form.include_seller_rollover and i + 1 < len(paragraphs):
                # Keep the content paragraph; strip highlight and enable auto-numbering
                content_para = paragraphs[i + 1]
                # Remove numId=0 so it inherits auto-numbering from StandardL1
                p_pr = content_para.find(_qn("w:pPr"))
                if p_pr is not None:
                    num_pr = p_pr.find(_qn("w:numPr"))
                    if num_pr is not None:
                        num_id = num_pr.find(_qn("w:numId"))
                        if num_id is not None and num_id.get(_qn("w:val")) == "0":
                            p_pr.remove(num_pr)
                    # Fix indentation to match other sections (left=0, firstLine=720)
                    ind = p_pr.find(_qn("w:ind"))
                    if ind is not None:
                        ind.set(_qn("w:left"), "0")
                        ind.set(_qn("w:firstLine"), "720")
                # Strip yellow highlighting from content runs
                for run in content_para.iterchildren(_qn("w:r")):
                    rpr = run.find(_qn("w:rPr"))
                    if rpr is not None:
                        hl = rpr.find(_qn("w:highlight"))
                        if hl is not None:
                            rpr.remove(hl)
            elif i + 1 < len(paragraphs):
                # Not included — delete the content paragraph too
                self._delete_entire_paragraph(paragraphs[i + 1], now)
            return

    def _handle_signature_block(self, body, form: LoiFormData, now: str):
        marker = "USE THE FOLLOWING SIGNATURE BLOCK STRUCTURE FOR PARCELS THAT ARE NOT OWNED BY INDIVIDUALS"
        seller_name_marker = "[SELLER NAME]"
        company_name_marker = "[COMPANY NAME]"

        paragraphs = list(body.iterchildren(_qn("w:p")))
        marker_idx = -1
        seller_name_idx = -1
        company_name_idx = -1

        for i, para in enumerate(paragraphs):
            text = _get_paragraph_text(para)
            if marker in text:
                marker_idx = i
            if seller_name_marker in text:
                seller_name_idx = i
            if company_name_marker in text:
                company_name_idx = i

        if marker_idx < 0:
            return

        if form.signature_block_type == SignatureBlockType.INDIVIDUAL:
            for i in range(marker_idx, len(paragraphs)):
                text = _get_paragraph_text(paragraphs[i])
                if "EXHIBIT A" in text or "WHEN A SINGLE LOI" in text:
                    break
                self._delete_entire_paragraph(paragraphs[i], now)
        else:
            # Delete individual seller signature line and name
            if seller_name_idx >= 0:
                self._delete_entire_paragraph(paragraphs[seller_name_idx], now)
                if seller_name_idx > 0:
                    prev_text = _get_paragraph_text(paragraphs[seller_name_idx - 1])
                    if "____" in prev_text:
                        self._delete_entire_paragraph(paragraphs[seller_name_idx - 1], now)

            # Delete the marker instruction
            self._delete_entire_paragraph(paragraphs[marker_idx], now)

            # Clone company block for multiple entities
            if company_name_idx >= 0 and len(form.signature_entities) > 1:
                block_start = company_name_idx
                block_end = company_name_idx
                for i in range(company_name_idx + 1, len(paragraphs)):
                    text = _get_paragraph_text(paragraphs[i])
                    if "WHEN A SINGLE LOI" in text or "EXHIBIT A" in text:
                        break
                    block_end = i
                    if text.startswith("Title:") or "Title:\t" in text or "Title:_" in text:
                        break

                insert_after = paragraphs[block_end]
                for e in range(1, len(form.signature_entities)):
                    spacer = _make_empty_paragraph()
                    insert_after.addnext(spacer)
                    insert_after = spacer

                    for p in range(block_start, block_end + 1):
                        clone = copy.deepcopy(paragraphs[p])
                        insert_after.addnext(clone)
                        insert_after = clone

    def _handle_multiple_entities(self, body, form: LoiFormData, now: str):
        marker = "WHEN A SINGLE LOI INCLUDES MULTIPLE ENTITIES"
        for para in list(body.iterchildren(_qn("w:p"))):
            text = _get_paragraph_text(para)
            if marker in text:
                self._delete_entire_paragraph(para, now)
                break

        if form.signature_block_type == SignatureBlockType.COMPANY_ENTITY and len(form.signature_entities) > 1:
            for para in list(body.iterchildren(_qn("w:p"))):
                text = _get_paragraph_text(para)
                if "(\u201cSeller\u201d)" in text:
                    target = "(\u201cSeller\u201d)"
                    replacement = "(collectively, the \u201cSeller\u201d)"
                    self._replace_text_in_paragraph(para, target, replacement, now)
                    break
                elif '("Seller")' in text:
                    self._replace_text_in_paragraph(para, '("Seller")', '(collectively, the "Seller")', now)
                    break

    def _handle_optional_paragraph(self, body, marker: str, include: bool, now: str):
        for para in list(body.iterchildren(_qn("w:p"))):
            text = _get_paragraph_text(para)
            if marker not in text:
                continue
            if include:
                marker_idx = text.index(marker)
                segments = [
                    (marker_idx, len(marker), True),
                    (0, marker_idx, False),
                    (marker_idx + len(marker), len(text) - marker_idx - len(marker), False),
                ]
                self._rebuild_paragraph_with_scenario(para, text, segments, now)
                # Remove numId=0 override so the paragraph inherits auto-numbering from its style
                p_pr = para.find(_qn("w:pPr"))
                if p_pr is not None:
                    num_pr = p_pr.find(_qn("w:numPr"))
                    if num_pr is not None:
                        num_id = num_pr.find(_qn("w:numId"))
                        if num_id is not None and num_id.get(_qn("w:val")) == "0":
                            p_pr.remove(num_pr)
            else:
                self._delete_entire_paragraph(para, now)
            return

    # ------------------------------------------------------------------
    # Placeholder Replacement
    # ------------------------------------------------------------------
    def _replace_placeholders(self, body, form: LoiFormData, now: str):
        # Seller address lines — replace positionally (empty lines delete the paragraph)
        self._replace_address_lines(body, "[____________________]",
            [form.seller_address_line1, form.seller_address_line2, form.seller_address_line3], now)

        # Deposit amounts ($10K placeholder)
        # Only include AdditionalDeposit when the selected scenario actually uses it
        deposit_values = []
        if form.initial_deposit is not None:
            deposit_values.append(to_legal_dollar_string(form.initial_deposit))
        if form.additional_deposit is not None and form.deposit_structure in (
            DepositStructure.GOVERNMENTAL_APPROVALS_GOING_HARD,
            DepositStructure.DUE_DILIGENCE_GOING_HARD,
            DepositStructure.MONTHLY_GOING_HARD,
        ):
            deposit_values.append(to_legal_dollar_string(form.additional_deposit))
        self._replace_sequential(body, "[Ten Thousand and 00/100 Dollars ($10,000.00)]", deposit_values, now)

        # $5K amounts (monthly release, legal reimb, extension deposit)
        five_k_values = []
        if form.deposit_structure == DepositStructure.MONTHLY_GOING_HARD and form.monthly_release_amount is not None:
            five_k_values.append(to_legal_dollar_string(form.monthly_release_amount))
        if form.include_legal_reimbursement and form.legal_reimbursement_amount is not None:
            five_k_values.append(to_legal_dollar_string(form.legal_reimbursement_amount))
        if form.include_option_to_extend and form.extension_deposit_amount is not None:
            five_k_values.append(to_legal_dollar_string(form.extension_deposit_amount))
        self._replace_sequential(body, "[Five Thousand and 00/100 Dollars ($5,000.00)]", five_k_values, now)

        # Company signature blocks
        if form.signature_block_type == SignatureBlockType.COMPANY_ENTITY:
            company_names = [e.company_name for e in form.signature_entities if e.company_name]
            self._replace_sequential(body, "[COMPANY NAME]", company_names, now)

        # All other replacements
        replacements = self._build_replacement_map(form)
        for para in list(body.iterchildren(_qn("w:p"))):
            for search, replace in replacements:
                text = _get_paragraph_text(para)
                if search in text:
                    self._replace_text_in_paragraph(para, search, replace, now)

        # Parcel IDs — targeted Exhibit A search (avoids matching stray "[]" elsewhere)
        self._replace_parcel_ids(body, form, now)

    def _replace_sequential(self, body, placeholder: str, values: list, now: str):
        if not values:
            return
        idx = 0
        for para in list(body.iterchildren(_qn("w:p"))):
            if idx >= len(values):
                break
            text = _get_paragraph_text(para)
            while placeholder in text and idx < len(values):
                self._replace_text_in_paragraph(para, placeholder, values[idx], now)
                idx += 1
                text = _get_paragraph_text(para)

    def _replace_address_lines(self, body, placeholder: str, values: list, now: str):
        """Replace address line placeholders positionally (1:1 mapping).
        Empty lines get their paragraph deleted instead of collapsing upward."""
        idx = 0
        for para in list(body.iterchildren(_qn("w:p"))):
            if idx >= len(values):
                break
            text = _get_paragraph_text(para)
            if placeholder not in text:
                continue
            if not values[idx]:
                self._delete_entire_paragraph(para, now)
            else:
                self._replace_text_in_paragraph(para, placeholder, values[idx], now)
            idx += 1

    def _replace_parcel_ids(self, body, form: LoiFormData, now: str):
        """Replace parcel ID placeholder only within the Exhibit A section,
        avoiding false matches of '[]' elsewhere in the document."""
        parcel_text = "\n".join(p for p in form.parcel_ids if p.strip())
        if not parcel_text:
            return
        in_exhibit_a = False
        for para in list(body.iterchildren(_qn("w:p"))):
            text = _get_paragraph_text(para)
            if "EXHIBIT A" in text or "Exhibit A" in text:
                in_exhibit_a = True
            if in_exhibit_a and "[]" in text:
                self._replace_text_in_paragraph(para, "[]", parcel_text, now)
                return

    def _build_replacement_map(self, form: LoiFormData) -> list:
        m = [("[Date]", form.date)]

        if form.attention_name:
            m.append(("[_______________]", form.attention_name))

        m.append(("[Address, City, State]", form.property_address))
        m.append(("[Mr./Mrs./Ms._________]", form.salutation))

        if form.seller_name:
            m.append(("[______________]", form.seller_name))

        if form.purchase_price is not None:
            pp = form.purchase_price
            words = convert_to_words(int(pp))
            number = f"{int(pp):,}"
            m.append(("[_________________]", words))
            m.append(("[_________]", number))

        if form.include_closing_extension and form.monthly_closing_extension_deposit is not None:
            m.append((
                "[Twenty-Five Thousand and 00/100 Dollars ($25,000.00)]",
                to_legal_dollar_string(form.monthly_closing_extension_deposit),
            ))

        if form.due_diligence_days is not None:
            m.append(("[one hundred twenty (120)]", _format_period(form.due_diligence_days)))
        if form.governmental_approvals_days is not None:
            m.append(("[one hundred fifty (150)]", _format_period(form.governmental_approvals_days)))
        if form.due_diligence_type == DueDiligenceType.WITH_ASSEMBLAGE and form.assemblage_days is not None:
            m.append(("[ninety (90)]", _format_period(form.assemblage_days)))
        if form.closing_days is not None:
            m.append(("[thirty (30)]", _format_period(form.closing_days)))
        if form.include_closing_extension and form.closing_extension_months is not None:
            m.append(("[six (6)]", _format_period(form.closing_extension_months)))

        if form.include_existing_leases:
            m.append(("[May 31, 2026]", form.lease_end_date))

        if form.include_lease_termination and form.lease_termination_days is not None:
            m.append(("[sixty (60)]", _format_period(form.lease_termination_days)))

        if form.commission_type == CommissionType.SUBTEXT_PAYS and form.broker_name:
            m.append(("[_____________]", form.broker_name))

        if form.signature_block_type == SignatureBlockType.INDIVIDUAL:
            m.append(("[SELLER NAME]", form.seller_name_signature))

        # Parcel IDs are handled separately in _replace_parcel_ids (targeted Exhibit A search)

        return m

    # ------------------------------------------------------------------
    # Header replacement
    # ------------------------------------------------------------------
    def _replace_header_placeholders(self, doc, form: LoiFormData, now: str):
        for section in doc.sections:
            header = section.header
            if header is None:
                continue
            for para in header._element.iterchildren(_qn("w:p")):
                text = _get_paragraph_text(para)
                if "[ADDRESS OR STREET NAME]" in text and form.property_address:
                    self._replace_text_in_paragraph(para, "[ADDRESS OR STREET NAME]", form.property_address, now)
                if "[DATE]" in text and form.date:
                    self._replace_text_in_paragraph(para, "[DATE]", form.date, now)

    # ------------------------------------------------------------------
    # Text replacement with tracked changes
    # ------------------------------------------------------------------
    def _replace_text_in_paragraph(self, para, search_text: str, replace_text: str, now: str):
        if not replace_text:
            return

        runs = list(para.iterchildren(_qn("w:r")))
        if not runs:
            return

        # Build character map
        char_map = []
        full_text_parts = []
        for run in runs:
            run_text = _get_run_text(run)
            for i, ch in enumerate(run_text):
                char_map.append((run, i))
                full_text_parts.append(ch)

        full_text = "".join(full_text_parts)
        match_start = full_text.find(search_text)
        if match_start < 0:
            return
        match_end = match_start + len(search_text)

        # Get run properties from the matched text
        match_rpr = _get_run_properties(char_map[match_start][0])

        # Calculate per-run positions
        cum_pos = 0
        run_ranges = []
        for r in runs:
            rlen = len(_get_run_text(r))
            run_ranges.append((r, cum_pos, cum_pos + rlen))
            cum_pos += rlen

        # Find affected runs
        affected = [(r, s, e) for r, s, e in run_ranges if s < match_end and e > match_start]
        if not affected:
            return

        first_run = affected[0][0]
        first_start = affected[0][1]
        last_run = affected[-1][0]
        last_end = affected[-1][2]

        new_elements = []

        # 1. Before-text fragment
        if first_start < match_start:
            before_text = _get_run_text(first_run)[:match_start - first_start]
            before_rpr = _get_run_properties(first_run)
            new_elements.append(_create_run(before_text, before_rpr))

        # 2. DeletedRun
        del_run = _make_deleted_run(search_text, match_rpr, self._author, now, self._next_rev_id())
        new_elements.append(del_run)

        # 3. InsertedRun
        ins_run = _make_inserted_run(replace_text, match_rpr, self._author, now, self._next_rev_id())
        new_elements.append(ins_run)

        # 4. After-text fragment
        if last_end > match_end:
            after_offset = match_end - affected[-1][1]
            after_text = _get_run_text(last_run)[after_offset:]
            after_rpr = _get_run_properties(last_run)
            new_elements.append(_create_run(after_text, after_rpr))

        # Insert before first affected run, then remove affected runs
        for elem in new_elements:
            first_run.addprevious(elem)

        for r, _, _ in affected:
            r.getparent().remove(r)

    # ------------------------------------------------------------------
    # Paragraph rebuild helpers
    # ------------------------------------------------------------------
    def _rebuild_paragraph_with_scenario(self, para, full_text: str,
                                          segments: list, now: str):
        runs = list(para.iterchildren(_qn("w:r")))
        if not runs:
            return

        # Build character -> RunProperties map
        char_props = []
        for run in runs:
            run_text = _get_run_text(run)
            rpr = _get_run_properties(run)
            for _ in run_text:
                char_props.append(rpr)

        # Preserve paragraph properties
        p_pr = para.find(_qn("w:pPr"))
        p_pr_copy = copy.deepcopy(p_pr) if p_pr is not None else None

        # Remove all existing runs
        for run in runs:
            para.remove(run)

        # Re-add paragraph properties if they were removed
        if p_pr_copy is not None and para.find(_qn("w:pPr")) is None:
            para.insert(0, p_pr_copy)

        # Sort segments by start position
        sorted_segments = sorted(segments, key=lambda s: s[0])

        for start, length, delete in sorted_segments:
            if length <= 0:
                continue
            actual_len = min(length, len(full_text) - start)
            segment_text = full_text[start:start + actual_len]
            if not segment_text:
                continue

            sub_runs = _split_by_formatting(segment_text, char_props, start)

            if delete:
                del_elem = etree.SubElement(para, _qn("w:del"))
                del_elem.set(_qn("w:author"), self._author)
                del_elem.set(_qn("w:date"), now)
                del_elem.set(_qn("w:id"), self._next_rev_id())
                for text, props in sub_runs:
                    run_elem = etree.SubElement(del_elem, _qn("w:r"))
                    if props is not None:
                        run_elem.append(copy.deepcopy(props))
                    dt = etree.SubElement(run_elem, _qn("w:delText"))
                    dt.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    dt.text = text
            else:
                for text, props in sub_runs:
                    run_elem = etree.SubElement(para, _qn("w:r"))
                    if props is not None:
                        props_copy = copy.deepcopy(props)
                        # Strip yellow highlighting from kept runs (template instructions use it)
                        hl = props_copy.find(_qn("w:highlight"))
                        if hl is not None:
                            props_copy.remove(hl)
                        run_elem.append(props_copy)
                    t = etree.SubElement(run_elem, _qn("w:t"))
                    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    t.text = text

    def _delete_entire_paragraph(self, para, now: str):
        text = _get_paragraph_text(para)
        if not text:
            return

        runs = list(para.iterchildren(_qn("w:r")))
        run_props = _get_run_properties(runs[0]) if runs else None

        for run in runs:
            para.remove(run)

        del_elem = etree.SubElement(para, _qn("w:del"))
        del_elem.set(_qn("w:author"), self._author)
        del_elem.set(_qn("w:date"), now)
        del_elem.set(_qn("w:id"), self._next_rev_id())

        run_elem = etree.SubElement(del_elem, _qn("w:r"))
        if run_props is not None:
            run_elem.append(copy.deepcopy(run_props))
        dt = etree.SubElement(run_elem, _qn("w:delText"))
        dt.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        dt.text = text

    # ------------------------------------------------------------------
    # Photo insertion
    # ------------------------------------------------------------------
    def _insert_property_photo(self, doc, body, form: LoiFormData, now: str):
        paragraphs = list(body.iterchildren(_qn("w:p")))

        # Find the parcel ID paragraph
        target_para = None
        for i in range(len(paragraphs) - 1, -1, -1):
            text = _get_all_paragraph_text(paragraphs[i])
            if any(p.strip() and p in text for p in form.parcel_ids):
                target_para = paragraphs[i]
                break

        # Fallback: look for "PARCEL ID" heading
        if target_para is None:
            for i in range(len(paragraphs) - 1, -1, -1):
                text = _get_all_paragraph_text(paragraphs[i])
                if "PARCEL ID" in text:
                    target_para = paragraphs[i + 2] if i + 2 < len(paragraphs) else paragraphs[i]
                    break

        # Final fallback
        if target_para is None:
            target_para = paragraphs[-1] if paragraphs else None
        if target_para is None:
            return

        # Get image dimensions using Pillow
        from PIL import Image as PILImage
        img = PILImage.open(io.BytesIO(form.property_photo_bytes))
        px_width, px_height = img.size

        # Convert to EMUs and scale
        max_width_emu = 5_040_000   # ~5.5 inches
        max_height_emu = 4_572_000  # ~5 inches

        img_width_emu = int(px_width * 914400 / 96)
        img_height_emu = int(px_height * 914400 / 96)

        if img_width_emu > max_width_emu:
            ratio = max_width_emu / img_width_emu
            img_width_emu = max_width_emu
            img_height_emu = int(img_height_emu * ratio)
        if img_height_emu > max_height_emu:
            ratio = max_height_emu / img_height_emu
            img_height_emu = max_height_emu
            img_width_emu = int(img_width_emu * ratio)

        # Determine content type for the image part
        content_type = form.property_photo_content_type or "image/jpeg"
        if "png" in content_type:
            img_type = "image/png"
        elif "gif" in content_type:
            img_type = "image/gif"
        elif "bmp" in content_type:
            img_type = "image/bmp"
        else:
            img_type = "image/jpeg"

        # Add image as a related part using OPC layer
        from docx.opc.part import Part
        from docx.opc.constants import RELATIONSHIP_TYPE as RT

        image_part_name = "/word/media/property_photo" + os.path.splitext(
            form.property_photo_filename or "photo.jpg")[1]

        # Create the image part
        from docx.opc.package import OpcPackage
        from docx.opc.packuri import PackURI

        part_name = PackURI(image_part_name)
        image_part = Part(part_name, img_type, form.property_photo_bytes, None)
        rel_id = doc.part.relate_to(image_part, RT.IMAGE)

        filename = form.property_photo_filename or "photo.jpg"

        # Build the inline image XML
        inline_xml = (
            f'<wp:inline distT="0" distB="0" distL="0" distR="0" '
            f'xmlns:wp="{WP_NS}" xmlns:a="{A_NS}" xmlns:pic="{PIC_NS}" xmlns:r="{R_NS}">'
            f'  <wp:extent cx="{img_width_emu}" cy="{img_height_emu}"/>'
            f'  <wp:effectExtent l="0" t="0" r="0" b="0"/>'
            f'  <wp:docPr id="1" name="Property Photo"/>'
            f'  <wp:cNvGraphicFramePr>'
            f'    <a:graphicFrameLocks noChangeAspect="1"/>'
            f'  </wp:cNvGraphicFramePr>'
            f'  <a:graphic>'
            f'    <a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
            f'      <pic:pic>'
            f'        <pic:nvPicPr>'
            f'          <pic:cNvPr id="0" name="{filename}"/>'
            f'          <pic:cNvPicPr/>'
            f'        </pic:nvPicPr>'
            f'        <pic:blipFill>'
            f'          <a:blip r:embed="{rel_id}"/>'
            f'          <a:stretch><a:fillRect/></a:stretch>'
            f'        </pic:blipFill>'
            f'        <pic:spPr>'
            f'          <a:xfrm>'
            f'            <a:off x="0" y="0"/>'
            f'            <a:ext cx="{img_width_emu}" cy="{img_height_emu}"/>'
            f'          </a:xfrm>'
            f'          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
            f'        </pic:spPr>'
            f'      </pic:pic>'
            f'    </a:graphicData>'
            f'  </a:graphic>'
            f'</wp:inline>'
        )
        inline_elem = etree.fromstring(inline_xml)

        # Create blank paragraphs
        blank1 = _make_empty_paragraph()
        blank2 = _make_empty_paragraph()

        # Create centered image paragraph
        img_para = etree.Element(_qn("w:p"))
        ppr = etree.SubElement(img_para, _qn("w:pPr"))
        jc = etree.SubElement(ppr, _qn("w:jc"))
        jc.set(_qn("w:val"), "center")

        run_elem = etree.SubElement(img_para, _qn("w:r"))
        drawing = etree.SubElement(run_elem, _qn("w:drawing"))
        drawing.append(inline_elem)

        # Insert after target
        target_para.addnext(blank1)
        blank1.addnext(blank2)
        blank2.addnext(img_para)


# ======================================================================
# Module-level helper functions
# ======================================================================

def _get_paragraph_text(para) -> str:
    """Get visible text from a paragraph's w:r elements (excludes tracked changes)."""
    texts = []
    for run in para.iterchildren(_qn("w:r")):
        texts.append(_get_run_text(run))
    return "".join(texts)


def _get_all_paragraph_text(para) -> str:
    """Get ALL text from a paragraph including tracked changes."""
    texts = []
    for t in para.iter(_qn("w:t")):
        texts.append(t.text or "")
    for dt in para.iter(_qn("w:delText")):
        texts.append(dt.text or "")
    return "".join(texts)


def _get_run_text(run) -> str:
    """Get text from a single run element."""
    texts = []
    for t in run.iterchildren(_qn("w:t")):
        texts.append(t.text or "")
    for dt in run.iterchildren(_qn("w:delText")):
        texts.append(dt.text or "")
    return "".join(texts)


def _get_run_properties(run):
    """Get the w:rPr element from a run, or None."""
    rpr = run.find(_qn("w:rPr"))
    return copy.deepcopy(rpr) if rpr is not None else None


def _create_run(text: str, rpr=None):
    """Create a w:r element with text."""
    run = etree.Element(_qn("w:r"))
    if rpr is not None:
        run.append(copy.deepcopy(rpr))
    t = etree.SubElement(run, _qn("w:t"))
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    return run


def _create_run_with_breaks(text: str, rpr=None):
    """Create a run with line breaks for multi-line text."""
    run = etree.Element(_qn("w:r"))
    if rpr is not None:
        run.append(copy.deepcopy(rpr))
    lines = text.split("\n")
    for i, line in enumerate(lines):
        if i > 0:
            etree.SubElement(run, _qn("w:br"))
        t = etree.SubElement(run, _qn("w:t"))
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = line
    return run


def _make_deleted_run(text: str, rpr, author: str, now: str, rev_id: str):
    """Create a w:del element containing the deleted text."""
    del_elem = etree.Element(_qn("w:del"))
    del_elem.set(_qn("w:author"), author)
    del_elem.set(_qn("w:date"), now)
    del_elem.set(_qn("w:id"), rev_id)

    run = etree.SubElement(del_elem, _qn("w:r"))
    if rpr is not None:
        run.append(copy.deepcopy(rpr))
    dt = etree.SubElement(run, _qn("w:delText"))
    dt.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    dt.text = text
    return del_elem


def _make_inserted_run(text: str, rpr, author: str, now: str, rev_id: str):
    """Create a w:ins element containing the inserted text."""
    ins_elem = etree.Element(_qn("w:ins"))
    ins_elem.set(_qn("w:author"), author)
    ins_elem.set(_qn("w:date"), now)
    ins_elem.set(_qn("w:id"), rev_id)

    if "\n" in text:
        run = _create_run_with_breaks(text, rpr)
    else:
        run = _create_run(text, rpr)
    ins_elem.append(run)
    return ins_elem


def _make_empty_paragraph():
    """Create an empty w:p element."""
    para = etree.Element(_qn("w:p"))
    run = etree.SubElement(para, _qn("w:r"))
    t = etree.SubElement(run, _qn("w:t"))
    t.text = ""
    return para


def _split_by_formatting(text: str, char_props: list, start_offset: int) -> list:
    """Split text into sub-runs based on formatting changes."""
    result = []
    if not text:
        return result

    current_props = char_props[start_offset] if start_offset < len(char_props) else None
    seg_start = 0

    for i in range(1, len(text)):
        char_idx = start_offset + i
        props = char_props[char_idx] if char_idx < len(char_props) else None

        prev_xml = etree.tostring(current_props).decode() if current_props is not None else ""
        cur_xml = etree.tostring(props).decode() if props is not None else ""

        if prev_xml != cur_xml:
            result.append((text[seg_start:i], current_props))
            seg_start = i
            current_props = props

    result.append((text[seg_start:], current_props))
    return result


def _format_period(value: int) -> str:
    """Convert number to 'word (number)' format, e.g. 120 -> 'one hundred twenty (120)'."""
    words = convert_to_words(value).lower()
    return f"{words} ({value})"
