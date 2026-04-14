"""
Comprehensive formatting test suite for the LOI Generator.

Generates Word documents across all scenario combinations and inspects:
- Tracked changes structure (w:del / w:ins elements)
- Placeholder replacement completeness (no stale brackets)
- Paragraph ordering and section lettering
- Deposit scenario text accuracy
- Closing extension text for MtM vs Single
- Commission scenario text
- Lease scenario combinations
- Signature block structure (individual vs entity, multiple entities)
- Seller rollover inclusion/exclusion
- Legal reimbursement inclusion/exclusion
- Option to extend customization
- Address line handling (empty lines)
- Parcel ID replacement in Exhibit A
- Dollar amount formatting
- Period formatting
"""

import io
import os
import sys
import copy
from itertools import product

# Add project root to path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from lxml import etree
from docx import Document

from services.document_generator import DocumentGenerator
from services.loi_form_data import (
    LoiFormData,
    DepositStructure,
    DueDiligenceType,
    ClosingExtensionType,
    CommissionType,
    SignatureBlockType,
    SignatureEntity,
)
from services.number_to_words import to_legal_dollar_string, convert_to_words

# ── Namespaces ──────────────────────────────────────────────────────────
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

def _qn(tag):
    prefix, local = tag.split(":")
    ns = {"w": W_NS}
    return f"{{{ns[prefix]}}}{local}"


# ── Helpers ─────────────────────────────────────────────────────────────
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_DIR = os.path.dirname(SCRIPT_DIR)
TEMPLATE_PATH = os.path.join(PROJECT_DIR, "Templates", "LOI_Template.docx")

passed = 0
failed = 0
failures = []


def check(condition, label):
    global passed, failed
    if condition:
        passed += 1
        print(f"  PASS: {label}")
    else:
        failed += 1
        failures.append(label)
        print(f"  FAIL: {label}")


def get_body_text(doc):
    """Get all visible text from the document body (non-deleted)."""
    body = doc.element.body
    texts = []
    for para in body.iterchildren(_qn("w:p")):
        texts.append(get_paragraph_visible_text(para))
    return "\n".join(texts)


def get_all_text(doc):
    """Get ALL text including deleted text."""
    body = doc.element.body
    texts = []
    for para in body.iterchildren(_qn("w:p")):
        para_texts = []
        for t in para.iter(_qn("w:t")):
            para_texts.append(t.text or "")
        for dt in para.iter(_qn("w:delText")):
            para_texts.append(dt.text or "")
        texts.append("".join(para_texts))
    return "\n".join(texts)


def get_paragraph_visible_text(para):
    """Get visible (non-deleted) text from a paragraph in document order."""
    texts = []
    for child in para:
        tag = child.tag.split("}")[1] if "}" in child.tag else child.tag
        if tag == "r":
            for t in child.iterchildren(_qn("w:t")):
                texts.append(t.text or "")
        elif tag == "ins":
            for run in child.iterchildren(_qn("w:r")):
                for t in run.iterchildren(_qn("w:t")):
                    texts.append(t.text or "")
        # Skip w:del elements (deleted text)
    return "".join(texts)


def get_paragraph_deleted_text(para):
    """Get deleted text from a paragraph."""
    texts = []
    for del_elem in para.iterchildren(_qn("w:del")):
        for run in del_elem.iterchildren(_qn("w:r")):
            for dt in run.iterchildren(_qn("w:delText")):
                texts.append(dt.text or "")
    return "".join(texts)


def get_paragraph_inserted_text(para):
    """Get inserted text from a paragraph."""
    texts = []
    for ins in para.iterchildren(_qn("w:ins")):
        for run in ins.iterchildren(_qn("w:r")):
            for t in run.iterchildren(_qn("w:t")):
                texts.append(t.text or "")
    return "".join(texts)


def get_paragraphs_text_list(doc):
    """Return list of (visible_text, deleted_text, inserted_text) per paragraph."""
    body = doc.element.body
    result = []
    for para in body.iterchildren(_qn("w:p")):
        vis = get_paragraph_visible_text(para)
        dele = get_paragraph_deleted_text(para)
        ins = get_paragraph_inserted_text(para)
        result.append((vis, dele, ins))
    return result


def generate(form):
    """Generate a document and return the Document object."""
    gen = DocumentGenerator()
    buf = gen.generate(TEMPLATE_PATH, form)
    return Document(io.BytesIO(buf.read()))


def make_form(**overrides):
    """Create a LoiFormData with sensible defaults and apply overrides."""
    defaults = dict(
        date="April 14, 2026",
        seller_address_line1="ABC Corporation",
        seller_address_line2="123 Main Street",
        seller_address_line3="St. Louis, MO 63101",
        attention_name="John Smith",
        property_address="456 Oak Avenue, St. Louis, MO 63102",
        salutation="Mr. Smith",
        seller_name="ABC Corporation",
        purchase_price=500000.0,
        initial_deposit=10000.0,
        additional_deposit=10000.0,
        monthly_release_amount=5000.0,
        legal_reimbursement_amount=5000.0,
        extension_deposit_amount=5000.0,
        monthly_closing_extension_deposit=25000.0,
        due_diligence_days=120,
        governmental_approvals_days=150,
        assemblage_days=90,
        closing_days=30,
        closing_extension_months=6,
        lease_end_date="May 31, 2026",
        lease_termination_days=60,
        broker_name="Jones Realty",
        seller_name_signature="John Doe",
        signature_entities=[SignatureEntity(company_name="ABC Corporation")],
        parcel_ids=["12-345-678"],
        deposit_structure=DepositStructure.GOVERNMENTAL_APPROVALS_GOING_HARD,
        include_legal_reimbursement=False,
        due_diligence_type=DueDiligenceType.STANDARD,
        closing_extension_type=ClosingExtensionType.NONE,
        commission_type=CommissionType.SELLER_PAYS_LISTING_AGENT,
        include_option_to_extend=True,
        num_extension_options=2,
        extension_option_days=60,
        include_existing_leases=True,
        include_delivered_vacant=False,
        include_lease_termination=False,
        include_right_to_negotiate_with_tenants=False,
        include_seller_rollover=False,
        signature_block_type=SignatureBlockType.INDIVIDUAL,
        prepared_by_first_name="Jake",
        prepared_by_last_name="Test",
    )
    defaults.update(overrides)
    return LoiFormData(**defaults)


# ═══════════════════════════════════════════════════════════════════════
# TEST GROUPS
# ═══════════════════════════════════════════════════════════════════════

def test_stale_placeholders():
    """No unreplaced template placeholders should remain in the visible text."""
    print("\n=== Stale Placeholder Checks ===")

    # Template placeholders that should be replaced
    stale_markers = [
        "[Date]",
        "[____________________]",
        "[_______________]",
        "[Address, City, State]",
        "[Mr./Mrs./Ms._________]",
        "[______________]",
        "[_________________]",
        "[_________]",
        "[one hundred twenty (120)]",
        "[one hundred fifty (150)]",
        "[thirty (30)]",
        "[SELLER NAME]",
        "[May 31, 2026]",
    ]

    # Instruction markers that should be deleted
    instruction_markers = [
        "REMOVE PARAGRAPH ABOVE AND USE THE FOLLOWING",
        "USE THE FOLLOWING PARAGRAPH IF MONEY IS GOING HARD",
        "REPLACE THIS ENTIRE PARAGRAPH WITH THE FOLLOWING",
        "INSERT THE FOLLOWING LANGUAGE IF WE ARE INCLUDING A CLOSING EXTENSION",
        "IF SELLER IS PAYING COMMISSION TO LISTING AGENT",
        "IF SUBTEXT IS PAYING THE COMMISSION",
        "IF NO BROKERS ARE INVOLVED IN THIS TRANSACTION",
        "REMOVE THE PRIOR SENTENCE AND USE THE FOLLOWING IF BEING DELIVERD VACANT",
        "REMOVE THE PRIOR SENTENCE AND ADD THE FOLLOWING IF LEASES",
        "USE THE FOLLOWING IF WE NEED THE RIGHT TO NEGOTIATE",
        "USE THE FOLLOWING SIGNATURE BLOCK STRUCTURE",
        "WHEN A SINGLE LOI INCLUDES MULTIPLE ENTITIES",
        "REMOVE UNLESS WE ARE PAYING A LEGAL REIMBURSEMENT",
        "REMOVE UNLESS WE ARE ALLOWING THE SELLER TO CONTRIBUTE",
    ]

    form = make_form()
    doc = generate(form)
    visible = get_body_text(doc)

    for marker in stale_markers:
        check(marker not in visible, f"No stale placeholder: {marker[:50]}")

    for marker in instruction_markers:
        check(marker not in visible, f"No instruction marker visible: {marker[:60]}")


def test_deposit_scenarios():
    """Test all three deposit scenarios produce correct text."""
    print("\n=== Deposit Scenario Tests ===")

    # Gov Approvals Going Hard
    form = make_form(deposit_structure=DepositStructure.GOVERNMENTAL_APPROVALS_GOING_HARD)
    doc = generate(form)
    text = get_body_text(doc)
    check("Initial Deposit" in text, "Gov Approvals: has Initial Deposit")
    check("Additional Deposit" in text, "Gov Approvals: has Additional Deposit")
    check("Earnest Money" in text, "Gov Approvals: has Earnest Money")
    check("Governmental Approvals Period" in text, "Gov Approvals: has Gov Approvals Period reference")
    # Should NOT contain Monthly Going Hard language
    check("Monthly Releases" not in text, "Gov Approvals: no Monthly Releases text")

    # DD Going Hard
    form = make_form(deposit_structure=DepositStructure.DUE_DILIGENCE_GOING_HARD)
    doc = generate(form)
    text = get_body_text(doc)
    check("Initial Deposit" in text, "DD Going Hard: has Initial Deposit")
    check("Additional Deposit" in text, "DD Going Hard: has Additional Deposit")
    check("waiver of the Due Diligence Period" in text, "DD Going Hard: has waiver language")
    check("Monthly Releases" not in text, "DD Going Hard: no Monthly Releases")

    # Monthly Going Hard
    form = make_form(deposit_structure=DepositStructure.MONTHLY_GOING_HARD,
                     monthly_release_amount=5000.0)
    doc = generate(form)
    text = get_body_text(doc)
    check("Initial Deposit" in text, "Monthly: has Initial Deposit")
    check("Additional Deposit" in text, "Monthly: has Additional Deposit")
    check("Monthly Releases" in text, "Monthly: has Monthly Releases language")
    check("Five Thousand" in text or "5,000" in text, "Monthly: has monthly release amount")


def test_deposit_dollar_amounts():
    """Verify deposit amounts are correctly formatted and placed."""
    print("\n=== Deposit Dollar Amount Tests ===")

    form = make_form(
        deposit_structure=DepositStructure.GOVERNMENTAL_APPROVALS_GOING_HARD,
        initial_deposit=15000.0,
        additional_deposit=20000.0,
    )
    doc = generate(form)
    text = get_body_text(doc)

    check("Fifteen Thousand and 00/100 Dollars ($15,000.00)" in text,
          "Gov: Initial deposit = $15K formatted correctly")
    check("Twenty Thousand and 00/100 Dollars ($20,000.00)" in text,
          "Gov: Additional deposit = $20K formatted correctly")

    # Monthly scenario with release amount
    form = make_form(
        deposit_structure=DepositStructure.MONTHLY_GOING_HARD,
        initial_deposit=25000.0,
        additional_deposit=25000.0,
        monthly_release_amount=7500.0,
    )
    doc = generate(form)
    text = get_body_text(doc)
    check("Twenty-Five Thousand and 00/100 Dollars ($25,000.00)" in text,
          "Monthly: $25K deposits formatted correctly")
    check("Seven Thousand Five Hundred and 00/100 Dollars ($7,500.00)" in text,
          "Monthly: monthly release $7.5K formatted correctly")


def test_legal_reimbursement():
    """Legal reimbursement paragraph appears/disappears correctly."""
    print("\n=== Legal Reimbursement Tests ===")

    # Without legal reimb
    form = make_form(include_legal_reimbursement=False)
    doc = generate(form)
    text = get_body_text(doc)
    check("Legal Reimbursement Fee" not in text, "No legal reimb: paragraph removed")

    # With legal reimb
    form = make_form(include_legal_reimbursement=True, legal_reimbursement_amount=7500.0)
    doc = generate(form)
    text = get_body_text(doc)
    check("Legal Reimbursement Fee" in text, "With legal reimb: paragraph present")
    check("Seven Thousand Five Hundred and 00/100 Dollars ($7,500.00)" in text,
          "Legal reimb: amount formatted correctly")


def test_due_diligence_scenarios():
    """Standard DD vs Assemblage produce different text."""
    print("\n=== Due Diligence Scenario Tests ===")

    # Standard
    form = make_form(due_diligence_type=DueDiligenceType.STANDARD,
                     due_diligence_days=100, governmental_approvals_days=180)
    doc = generate(form)
    text = get_body_text(doc)
    check("one hundred (100)" in text, "Standard DD: custom DD days")
    check("one hundred eighty (180)" in text, "Standard DD: custom GA days")
    check("Assemblage Period" not in text, "Standard DD: no assemblage text")

    # With Assemblage
    form = make_form(due_diligence_type=DueDiligenceType.WITH_ASSEMBLAGE,
                     due_diligence_days=120, governmental_approvals_days=150,
                     assemblage_days=90)
    doc = generate(form)
    text = get_body_text(doc)
    check("Assemblage Period" in text, "Assemblage DD: has assemblage text")
    check("ninety (90)" in text, "Assemblage DD: assemblage days present")


def test_closing_extension_scenarios():
    """All closing extension types produce correct output."""
    print("\n=== Closing Extension Tests ===")

    # No extension
    form = make_form(closing_extension_type=ClosingExtensionType.NONE)
    doc = generate(form)
    text = get_body_text(doc)
    check("Closing Extension" not in text, "No extension: no extension text")

    # Month-to-Month
    form = make_form(closing_extension_type=ClosingExtensionType.MONTH_TO_MONTH,
                     closing_extension_months=6,
                     monthly_closing_extension_deposit=25000.0)
    doc = generate(form)
    text = get_body_text(doc)
    check("month-to-month" in text, "MtM: has month-to-month language")
    check("Closing Extension" in text, "MtM: has Closing Extension")
    check("Monthly Closing Extension Deposit" in text, "MtM: has Monthly Closing Extension Deposit")

    # Single Extension
    form = make_form(closing_extension_type=ClosingExtensionType.SINGLE,
                     closing_extension_months=6,
                     monthly_closing_extension_deposit=25000.0)
    doc = generate(form)
    text = get_body_text(doc)
    check("Closing Extension" in text, "Single: has Closing Extension")
    check("Closing Extension Deposit" in text, "Single: has Closing Extension Deposit")
    # Should NOT have month-to-month language
    check("month-to-month" not in text, "Single: no month-to-month language")
    # Should have the "for a period of" language
    check("for a period of" in text, "Single: has 'for a period of' language")


def test_commission_scenarios():
    """All commission types produce correct text."""
    print("\n=== Commission Scenario Tests ===")

    # Seller pays
    form = make_form(commission_type=CommissionType.SELLER_PAYS_LISTING_AGENT)
    doc = generate(form)
    text = get_body_text(doc)
    check("Seller to pay the commission" in text, "Seller pays: correct text")
    check("listing agreement" in text, "Seller pays: listing agreement reference")

    # Subtext pays
    form = make_form(commission_type=CommissionType.SUBTEXT_PAYS, broker_name="Jones Realty")
    doc = generate(form)
    text = get_body_text(doc)
    check("Purchaser to pay a commission" in text, "Subtext pays: correct text")
    check("Jones Realty" in text, "Subtext pays: broker name present")

    # No brokers
    form = make_form(commission_type=CommissionType.NO_BROKERS)
    doc = generate(form)
    text = get_body_text(doc)
    check("no brokerage commission" in text, "No brokers: correct text")


def test_option_to_extend():
    """Option to extend: included with custom values or excluded."""
    print("\n=== Option to Extend Tests ===")

    # Included with defaults
    form = make_form(include_option_to_extend=True,
                     num_extension_options=2,
                     extension_option_days=60,
                     extension_deposit_amount=5000.0)
    doc = generate(form)
    text = get_body_text(doc)
    check("Option to Extend" in text, "Included: Option to Extend present")
    check("Extension Notice" in text, "Included: Extension Notice text")

    # With custom values
    form = make_form(include_option_to_extend=True,
                     num_extension_options=3,
                     extension_option_days=90,
                     extension_deposit_amount=10000.0)
    doc = generate(form)
    text = get_body_text(doc)
    check("three (3)" in text, "Custom: 3 extensions")
    check("ninety (90)" in text, "Custom: 90 day extensions")
    check("Ten Thousand and 00/100 Dollars ($10,000.00)" in text, "Custom: $10K deposit")

    # Excluded
    form = make_form(include_option_to_extend=False)
    doc = generate(form)
    text = get_body_text(doc)
    # The paragraph should be deleted (all text in w:del)
    paragraphs = get_paragraphs_text_list(doc)
    option_visible = any("Option to Extend" in vis for vis, _, _ in paragraphs)
    option_deleted = any("Option to Extend" in dele for _, dele, _ in paragraphs)
    check(not option_visible or option_deleted, "Excluded: Option to Extend paragraph deleted")


def test_lease_scenarios():
    """Various lease option combinations."""
    print("\n=== Lease Scenario Tests ===")

    # Existing leases only
    form = make_form(include_existing_leases=True,
                     include_delivered_vacant=False,
                     include_lease_termination=False,
                     include_right_to_negotiate_with_tenants=False)
    doc = generate(form)
    text = get_body_text(doc)
    check("existing leases" in text, "Existing leases: text present")
    check("May 31, 2026" in text, "Existing leases: end date present")

    # Delivered vacant
    form = make_form(include_existing_leases=False,
                     include_delivered_vacant=True,
                     include_lease_termination=False,
                     include_right_to_negotiate_with_tenants=False)
    doc = generate(form)
    text = get_body_text(doc)
    check("delivered vacant" in text, "Delivered vacant: text present")

    # Lease termination
    form = make_form(include_existing_leases=True,
                     include_delivered_vacant=False,
                     include_lease_termination=True,
                     lease_termination_days=90,
                     include_right_to_negotiate_with_tenants=False)
    doc = generate(form)
    text = get_body_text(doc)
    check("termination provision" in text, "Lease termination: text present")
    check("ninety (90)" in text, "Lease termination: 90 days present")

    # Right to negotiate
    form = make_form(include_existing_leases=True,
                     include_delivered_vacant=False,
                     include_lease_termination=False,
                     include_right_to_negotiate_with_tenants=True)
    doc = generate(form)
    text = get_body_text(doc)
    check("negotiate directly" in text, "Right to negotiate: text present")

    # All lease options ON
    form = make_form(include_existing_leases=True,
                     include_delivered_vacant=True,
                     include_lease_termination=True,
                     include_right_to_negotiate_with_tenants=True,
                     lease_termination_days=60)
    doc = generate(form)
    text = get_body_text(doc)
    check("existing leases" in text, "All leases: existing leases text")
    check("delivered vacant" in text or "terminate" in text, "All leases: vacant/termination text")
    check("negotiate directly" in text, "All leases: negotiate text")


def test_seller_rollover():
    """Seller rollover paragraph appears/disappears correctly."""
    print("\n=== Seller Rollover Tests ===")

    # Without rollover
    form = make_form(include_seller_rollover=False)
    doc = generate(form)
    text = get_body_text(doc)
    check("Seller Rollover" not in text, "No rollover: paragraph removed from visible text")

    # With rollover
    form = make_form(include_seller_rollover=True)
    doc = generate(form)
    text = get_body_text(doc)
    check("Seller Rollover" in text or "passive investor" in text,
          "With rollover: paragraph present")


def test_signature_block_individual():
    """Individual seller signature block."""
    print("\n=== Individual Signature Block Tests ===")

    form = make_form(
        signature_block_type=SignatureBlockType.INDIVIDUAL,
        seller_name_signature="Jane Doe",
    )
    doc = generate(form)
    text = get_body_text(doc)
    check("Jane Doe" in text, "Individual: seller name in signature")
    # Company block markers should be deleted
    check("[COMPANY NAME]" not in text, "Individual: no company name placeholder visible")


def test_signature_block_single_entity():
    """Single company entity signature block."""
    print("\n=== Single Entity Signature Block Tests ===")

    form = make_form(
        signature_block_type=SignatureBlockType.COMPANY_ENTITY,
        signature_entities=[SignatureEntity(company_name="XYZ Holdings LLC")],
        seller_name="XYZ Holdings LLC",
    )
    doc = generate(form)
    text = get_body_text(doc)
    check("XYZ Holdings LLC" in text, "Single entity: company name present")
    # Should NOT have "collectively, the Seller" (note: "collectively" appears
    # naturally in template text like "collectively, Governmental Approvals")
    check('collectively, the \u201cSeller\u201d' not in text and 'collectively, the "Seller"' not in text,
          "Single entity: no collectively-the-Seller language")


def test_signature_block_multiple_entities():
    """Multiple company entities signature block."""
    print("\n=== Multiple Entity Signature Block Tests ===")

    form = make_form(
        signature_block_type=SignatureBlockType.COMPANY_ENTITY,
        signature_entities=[
            SignatureEntity(company_name="Alpha LLC"),
            SignatureEntity(company_name="Beta Corp"),
        ],
        seller_name="Alpha LLC and Beta Corp",
    )
    doc = generate(form)
    text = get_body_text(doc)
    check("Alpha LLC" in text, "Multi entity: first company name present")
    check("Beta Corp" in text, "Multi entity: second company name present")
    check("collectively" in text, "Multi entity: collectively language present")


def test_address_lines():
    """Address lines: all present, empty middle, only first."""
    print("\n=== Address Line Tests ===")

    # All 3 lines
    form = make_form(
        seller_address_line1="ABC Corporation",
        seller_address_line2="123 Main St",
        seller_address_line3="St. Louis, MO 63101",
    )
    doc = generate(form)
    text = get_body_text(doc)
    check("ABC Corporation" in text, "All lines: line 1 present")
    check("123 Main St" in text, "All lines: line 2 present")
    check("St. Louis, MO 63101" in text, "All lines: line 3 present")
    check("[____________________]" not in text, "All lines: no stale placeholder")

    # Empty line 2
    form = make_form(
        seller_address_line1="ABC Corporation",
        seller_address_line2="",
        seller_address_line3="St. Louis, MO 63101",
    )
    doc = generate(form)
    text = get_body_text(doc)
    check("ABC Corporation" in text, "Empty L2: line 1 present")
    check("St. Louis, MO 63101" in text, "Empty L2: line 3 present")
    # Line 2's placeholder paragraph should be deleted entirely
    check("[____________________]" not in text, "Empty L2: no stale placeholder")

    # Only line 1
    form = make_form(
        seller_address_line1="ABC Corporation",
        seller_address_line2="",
        seller_address_line3="",
    )
    doc = generate(form)
    text = get_body_text(doc)
    check("ABC Corporation" in text, "Only L1: line 1 present")


def test_parcel_ids():
    """Parcel IDs replaced in Exhibit A section."""
    print("\n=== Parcel ID Tests ===")

    # Single parcel
    form = make_form(parcel_ids=["12-345-678"])
    doc = generate(form)
    text = get_body_text(doc)
    check("12-345-678" in text, "Single parcel: ID present")

    # Multiple parcels
    form = make_form(parcel_ids=["12-345-678", "98-765-432", "11-222-333"])
    doc = generate(form)
    text = get_body_text(doc)
    check("12-345-678" in text, "Multi parcel: ID 1 present")
    check("98-765-432" in text, "Multi parcel: ID 2 present")
    check("11-222-333" in text, "Multi parcel: ID 3 present")


def test_purchase_price_formatting():
    """Purchase price words and number formatting."""
    print("\n=== Purchase Price Formatting Tests ===")

    # $500,000
    form = make_form(purchase_price=500000.0)
    doc = generate(form)
    text = get_body_text(doc)
    check("Five Hundred Thousand" in text, "$500K: words present")
    check("500,000" in text, "$500K: number present")

    # $1,250,000
    form = make_form(purchase_price=1250000.0)
    doc = generate(form)
    text = get_body_text(doc)
    check("One Million Two Hundred Fifty Thousand" in text, "$1.25M: words present")
    check("1,250,000" in text, "$1.25M: number present")

    # $10,000,000
    form = make_form(purchase_price=10000000.0)
    doc = generate(form)
    text = get_body_text(doc)
    check("Ten Million" in text, "$10M: words present")
    check("10,000,000" in text, "$10M: number present")


def test_period_formatting():
    """Period values in 'word (number)' format."""
    print("\n=== Period Formatting Tests ===")

    form = make_form(
        due_diligence_days=45,
        governmental_approvals_days=200,
        closing_days=60,
    )
    doc = generate(form)
    text = get_body_text(doc)
    check("forty-five (45)" in text, "45 days formatted correctly")
    check("two hundred (200)" in text, "200 days formatted correctly")
    check("sixty (60)" in text, "60 days formatted correctly")


def test_tracked_changes_structure():
    """Every replacement should produce w:del + w:ins tracked change pairs."""
    print("\n=== Tracked Changes Structure Tests ===")

    form = make_form()
    doc = generate(form)
    body = doc.element.body

    # Count tracked change elements
    del_count = len(list(body.iter(_qn("w:del"))))
    ins_count = len(list(body.iter(_qn("w:ins"))))

    check(del_count > 0, f"Has w:del elements ({del_count})")
    check(ins_count > 0, f"Has w:ins elements ({ins_count})")

    # Author should be set
    for del_elem in body.iter(_qn("w:del")):
        author = del_elem.get(_qn("w:author"))
        check(author is not None and author != "", "w:del has author attribute")
        break  # Just check first one

    for ins_elem in body.iter(_qn("w:ins")):
        author = ins_elem.get(_qn("w:author"))
        check(author is not None and author != "", "w:ins has author attribute")
        break


def test_header_replacements():
    """Header placeholders are replaced."""
    print("\n=== Header Replacement Tests ===")

    form = make_form(property_address="789 Elm Street, Chicago, IL", date="March 15, 2026")
    doc = generate(form)

    for section in doc.sections:
        header = section.header
        if header is None:
            continue
        header_text = ""
        for para in header._element.iterchildren(_qn("w:p")):
            header_text += get_paragraph_visible_text(para) + " "
            # Also check inserted text in header
            header_text += get_paragraph_inserted_text(para) + " "

        if "[ADDRESS OR STREET NAME]" in header_text or "[DATE]" in header_text:
            check(False, "Header: stale placeholders remain")
        else:
            check(True, "Header: no stale placeholders")
        break  # Check first section only


def test_section_lettering_basic():
    """Section names appear in correct order; letters are auto-numbered by Word style.
    Deleted paragraphs with StandardL1 must have numId=0 to avoid consuming letters."""
    print("\n=== Section Lettering Tests ===")

    form = make_form(include_seller_rollover=False)
    doc = generate(form)
    body = doc.element.body

    # Expected section names in order (letters assigned by Word auto-numbering)
    expected_names = [
        "Property.",
        "Purchase Price.",
        "Deposit.",
        # Legal Reimbursement (deleted) — should have numId=0
        "Exclusivity",
        "Due Diligence",
        "Closing.",
        "Commissions.",
        "Option to Extend.",
        "Leases.",
        # Seller Rollover (deleted) — should have numId=0
        "Confidentiality.",
        "Governing Law.",
        "Miscellaneous.",
    ]

    found_order = []
    for para in body.iterchildren(_qn("w:p")):
        vis = get_paragraph_visible_text(para)
        for name in expected_names:
            if name in vis and name not in found_order:
                found_order.append(name)
                break

    for name in expected_names:
        check(name in found_order, f"Section '{name}' found in document")

    # Verify order is correct
    check(found_order == expected_names,
          f"Sections appear in correct order ({len(found_order)}/{len(expected_names)} found)")

    # Verify deleted paragraphs have numId=0 (don't consume letters)
    for para in body.iterchildren(_qn("w:p")):
        has_del = any(True for _ in para.iterchildren(_qn("w:del")))
        vis = get_paragraph_visible_text(para).strip()
        if has_del and not vis:
            p_pr = para.find(_qn("w:pPr"))
            if p_pr is not None:
                p_style = p_pr.find(_qn("w:pStyle"))
                if p_style is not None and "Standard" in (p_style.get(_qn("w:val")) or ""):
                    num_pr = p_pr.find(_qn("w:numPr"))
                    num_id_elem = num_pr.find(_qn("w:numId")) if num_pr is not None else None
                    num_id = num_id_elem.get(_qn("w:val")) if num_id_elem is not None else None
                    del_text = get_paragraph_deleted_text(para)[:50]
                    check(num_id == "0",
                          f"Deleted StandardL1 para has numId=0: {del_text}...")


def test_section_lettering_with_rollover():
    """When seller rollover is included, section letters should shift."""
    print("\n=== Section Lettering with Rollover ===")

    form = make_form(include_seller_rollover=True)
    doc = generate(form)
    text = get_body_text(doc)

    # With rollover, there should be an extra section between Leases and Confidentiality
    check("Seller Rollover" in text or "passive investor" in text,
          "Rollover: section present")


def test_closing_extension_single_text_accuracy():
    """Single closing extension should have exact expected wording."""
    print("\n=== Single Closing Extension Text Accuracy ===")

    form = make_form(
        closing_extension_type=ClosingExtensionType.SINGLE,
        closing_extension_months=3,
        monthly_closing_extension_deposit=50000.0,
    )
    doc = generate(form)
    text = get_body_text(doc)

    check("for a period of" in text, "Single ext: 'for a period of' present")
    check("three (3)" in text, "Single ext: 3 months formatted")
    check("Fifty Thousand and 00/100 Dollars ($50,000.00)" in text, "Single ext: $50K deposit")
    check("Closing Extension Deposit" in text, "Single ext: Closing Extension Deposit label")
    check("non-refundable" in text, "Single ext: non-refundable clause")
    # Single should have "the Closing Extension" not "each a Closing Extension"
    check("the \u201cClosing Extension\u201d" in text or 'the "Closing Extension"' in text,
          "Single ext: uses 'the' not 'each a'")


def test_closing_extension_mtm_text_accuracy():
    """Month-to-month closing extension should have exact expected wording."""
    print("\n=== MtM Closing Extension Text Accuracy ===")

    form = make_form(
        closing_extension_type=ClosingExtensionType.MONTH_TO_MONTH,
        closing_extension_months=8,
        monthly_closing_extension_deposit=30000.0,
    )
    doc = generate(form)
    text = get_body_text(doc)

    check("month-to-month" in text, "MtM ext: month-to-month present")
    check("eight (8)" in text, "MtM ext: 8 months formatted")
    check("Thirty Thousand and 00/100 Dollars ($30,000.00)" in text, "MtM ext: $30K deposit")
    check("Monthly Closing Extension Deposit" in text, "MtM ext: Monthly label")
    check("for each month" in text, "MtM ext: 'for each month' language")


def test_full_scenario_matrix_generates_without_error():
    """Generate docs for a matrix of scenario combinations — none should crash."""
    print("\n=== Full Scenario Matrix (no-crash) Tests ===")

    deposit_options = list(DepositStructure)
    dd_options = list(DueDiligenceType)
    closing_ext_options = list(ClosingExtensionType)
    commission_options = list(CommissionType)
    sig_options = list(SignatureBlockType)

    count = 0
    crash_count = 0

    for dep, dd, ext, comm, sig in product(
        deposit_options, dd_options, closing_ext_options, commission_options, sig_options
    ):
        count += 1
        try:
            form = make_form(
                deposit_structure=dep,
                due_diligence_type=dd,
                closing_extension_type=ext,
                commission_type=comm,
                signature_block_type=sig,
                broker_name="Test Broker" if comm == CommissionType.SUBTEXT_PAYS else "",
            )
            doc = generate(form)
            visible = get_body_text(doc)
            # Basic sanity: should have some text
            assert len(visible) > 100, f"Document too short: {len(visible)} chars"
        except Exception as e:
            crash_count += 1
            check(False, f"CRASH: dep={dep.name}, dd={dd.name}, ext={ext.name}, comm={comm.name}, sig={sig.name}: {e}")

    check(crash_count == 0, f"All {count} scenario combos generated without error ({crash_count} crashes)")


def test_lease_all_off():
    """When all lease options are off, the lease paragraph should still handle gracefully."""
    print("\n=== All Lease Options Off ===")

    form = make_form(
        include_existing_leases=False,
        include_delivered_vacant=False,
        include_lease_termination=False,
        include_right_to_negotiate_with_tenants=False,
    )
    try:
        doc = generate(form)
        text = get_body_text(doc)
        check(True, "All lease off: generates without error")
        # With all lease options off, "Leases." section header might still exist
        # but all sub-content should be deleted
    except Exception as e:
        check(False, f"All lease off: crashed: {e}")


def test_number_to_words_edge_cases():
    """Edge cases for number-to-words conversion."""
    print("\n=== Number to Words Edge Cases ===")

    check(convert_to_words(0) == "Zero", f"0 -> 'Zero' (got '{convert_to_words(0)}')")
    check(convert_to_words(1) == "One", f"1 -> 'One' (got '{convert_to_words(1)}')")
    check(convert_to_words(11) == "Eleven", f"11 -> 'Eleven' (got '{convert_to_words(11)}')")
    check(convert_to_words(100) == "One Hundred", f"100 -> 'One Hundred' (got '{convert_to_words(100)}')")
    check(convert_to_words(1000) == "One Thousand", f"1000 -> 'One Thousand' (got '{convert_to_words(1000)}')")
    check(convert_to_words(1000000) == "One Million", f"1M -> 'One Million' (got '{convert_to_words(1000000)}')")
    check(convert_to_words(999999) == "Nine Hundred Ninety-Nine Thousand Nine Hundred Ninety-Nine",
          f"999999 words (got '{convert_to_words(999999)}')")

    # Legal dollar string
    result = to_legal_dollar_string(1234567.0)
    check("One Million Two Hundred Thirty-Four Thousand Five Hundred Sixty-Seven" in result,
          f"$1,234,567 legal string words")
    check("$1,234,567.00" in result, f"$1,234,567 legal string number")


def test_dollar_format_in_legal_string():
    """Legal dollar strings should have correct format: 'Words and 00/100 Dollars ($N.00)'"""
    print("\n=== Legal Dollar String Format ===")

    tests = [
        (10000.0, "Ten Thousand and 00/100 Dollars ($10,000.00)"),
        (5000.0, "Five Thousand and 00/100 Dollars ($5,000.00)"),
        (25000.0, "Twenty-Five Thousand and 00/100 Dollars ($25,000.00)"),
        (100000.0, "One Hundred Thousand and 00/100 Dollars ($100,000.00)"),
        (1500000.0, "One Million Five Hundred Thousand and 00/100 Dollars ($1,500,000.00)"),
    ]
    for amount, expected in tests:
        result = to_legal_dollar_string(amount)
        check(result == expected, f"${amount:,.0f} -> '{expected}' (got '{result}')")


def test_deposit_scenario_no_cross_contamination():
    """Each deposit scenario should only contain its own text, not others'."""
    print("\n=== Deposit Cross-Contamination Tests ===")

    # Gov Approvals should NOT have "waiver of the Due Diligence Period" or "Monthly Releases"
    form = make_form(deposit_structure=DepositStructure.GOVERNMENTAL_APPROVALS_GOING_HARD)
    doc = generate(form)
    text = get_body_text(doc)
    check("waiver of the Due Diligence Period" not in text,
          "Gov: no DD waiver language contamination")
    check("Monthly Releases" not in text,
          "Gov: no Monthly Releases contamination")

    # DD Going Hard should NOT have "Monthly Releases"
    form = make_form(deposit_structure=DepositStructure.DUE_DILIGENCE_GOING_HARD)
    doc = generate(form)
    text = get_body_text(doc)
    check("Monthly Releases" not in text,
          "DD: no Monthly Releases contamination")

    # Monthly should NOT have the pure Gov Approvals "receipt of Governmental Approvals" as the sole refundability condition
    form = make_form(deposit_structure=DepositStructure.MONTHLY_GOING_HARD)
    doc = generate(form)
    text = get_body_text(doc)
    check("Monthly Releases" in text,
          "Monthly: has its own Monthly Releases language")


def test_five_k_placeholder_count():
    """$5K placeholders should match the number of $5K values provided."""
    print("\n=== $5K Placeholder Count Tests ===")

    # Monthly + legal reimb + option to extend = 3 five-K values
    form = make_form(
        deposit_structure=DepositStructure.MONTHLY_GOING_HARD,
        monthly_release_amount=5000.0,
        include_legal_reimbursement=True,
        legal_reimbursement_amount=7500.0,
        include_option_to_extend=True,
        extension_deposit_amount=8000.0,
    )
    doc = generate(form)
    text = get_body_text(doc)

    # Should contain all three amounts
    check("Five Thousand and 00/100 Dollars ($5,000.00)" in text, "$5K: monthly release present")
    check("Seven Thousand Five Hundred and 00/100 Dollars ($7,500.00)" in text, "$5K: legal reimb present")
    check("Eight Thousand and 00/100 Dollars ($8,000.00)" in text, "$5K: extension deposit present")

    # No stale $5K placeholder
    check("[Five Thousand and 00/100 Dollars ($5,000.00)]" not in text,
          "$5K: no stale placeholder remaining")


def test_ten_k_placeholder_count():
    """$10K placeholders should match the number of $10K values provided."""
    print("\n=== $10K Placeholder Count Tests ===")

    # Gov Approvals: initial + additional = 2 ten-K values
    form = make_form(
        deposit_structure=DepositStructure.GOVERNMENTAL_APPROVALS_GOING_HARD,
        initial_deposit=15000.0,
        additional_deposit=20000.0,
    )
    doc = generate(form)
    text = get_body_text(doc)
    check("[Ten Thousand and 00/100 Dollars ($10,000.00)]" not in text,
          "$10K gov: no stale placeholder remaining")

    # Monthly Going Hard: initial + additional = 2 ten-K values
    form = make_form(
        deposit_structure=DepositStructure.MONTHLY_GOING_HARD,
        initial_deposit=12000.0,
        additional_deposit=18000.0,
    )
    doc = generate(form)
    text = get_body_text(doc)
    check("[Ten Thousand and 00/100 Dollars ($10,000.00)]" not in text,
          "$10K monthly: no stale placeholder remaining")


def test_closing_extension_none_no_stale_placeholders():
    """When no closing extension, no extension placeholders should remain."""
    print("\n=== Closing Extension None: No Stale Placeholders ===")

    form = make_form(closing_extension_type=ClosingExtensionType.NONE)
    doc = generate(form)
    text = get_body_text(doc)
    check("[six (6)]" not in text, "No ext: no [six (6)] placeholder")
    check("[Twenty-Five Thousand" not in text, "No ext: no $25K placeholder")


def test_multiple_entities_collective_in_opening():
    """Multiple entities should trigger 'collectively, the Seller' in opening paragraph."""
    print("\n=== Multiple Entities Collective Language ===")

    form = make_form(
        signature_block_type=SignatureBlockType.COMPANY_ENTITY,
        signature_entities=[
            SignatureEntity(company_name="Alpha LLC"),
            SignatureEntity(company_name="Beta Corp"),
            SignatureEntity(company_name="Gamma Inc"),
        ],
        seller_name="Alpha LLC, Beta Corp and Gamma Inc",
    )
    doc = generate(form)
    text = get_body_text(doc)
    check("collectively" in text, "3 entities: collectively language present")


def test_empty_optional_fields():
    """Document should generate cleanly even with minimal form data."""
    print("\n=== Minimal Form Data Test ===")

    form = make_form(
        seller_address_line1="",
        seller_address_line2="",
        seller_address_line3="",
        attention_name="",
        salutation="",
    )
    try:
        doc = generate(form)
        check(True, "Minimal form: generates without error")
    except Exception as e:
        check(False, f"Minimal form: crashed: {e}")


def test_closing_days_replacement():
    """Closing days placeholder [thirty (30)] should be replaced."""
    print("\n=== Closing Days Replacement ===")

    form = make_form(closing_days=45)
    doc = generate(form)
    text = get_body_text(doc)
    check("forty-five (45)" in text, "Closing: 45 days formatted")
    check("[thirty (30)]" not in text, "Closing: no stale placeholder")


def test_lease_termination_days_replacement():
    """Lease termination days placeholder should be replaced when enabled."""
    print("\n=== Lease Termination Days ===")

    form = make_form(
        include_lease_termination=True,
        lease_termination_days=45,
    )
    doc = generate(form)
    text = get_body_text(doc)
    check("forty-five (45)" in text, "Lease term: 45 days formatted")


# ═══════════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════════

if __name__ == "__main__":
    print("=" * 70)
    print("  COMPREHENSIVE LOI FORMATTING TEST SUITE")
    print("=" * 70)

    if not os.path.exists(TEMPLATE_PATH):
        print(f"\nERROR: Template not found at {TEMPLATE_PATH}")
        sys.exit(1)

    test_stale_placeholders()
    test_deposit_scenarios()
    test_deposit_dollar_amounts()
    test_legal_reimbursement()
    test_due_diligence_scenarios()
    test_closing_extension_scenarios()
    test_commission_scenarios()
    test_option_to_extend()
    test_lease_scenarios()
    test_seller_rollover()
    test_signature_block_individual()
    test_signature_block_single_entity()
    test_signature_block_multiple_entities()
    test_address_lines()
    test_parcel_ids()
    test_purchase_price_formatting()
    test_period_formatting()
    test_tracked_changes_structure()
    test_header_replacements()
    test_section_lettering_basic()
    test_section_lettering_with_rollover()
    test_closing_extension_single_text_accuracy()
    test_closing_extension_mtm_text_accuracy()
    test_full_scenario_matrix_generates_without_error()
    test_lease_all_off()
    test_number_to_words_edge_cases()
    test_dollar_format_in_legal_string()
    test_deposit_scenario_no_cross_contamination()
    test_five_k_placeholder_count()
    test_ten_k_placeholder_count()
    test_closing_extension_none_no_stale_placeholders()
    test_multiple_entities_collective_in_opening()
    test_empty_optional_fields()
    test_closing_days_replacement()
    test_lease_termination_days_replacement()

    print("\n" + "=" * 70)
    print(f"  RESULTS: {passed} passed, {failed} failed")
    print("=" * 70)

    if failures:
        print("\n  FAILURES:")
        for f in failures:
            print(f"    - {f}")

    print()
    sys.exit(1 if failed > 0 else 0)
